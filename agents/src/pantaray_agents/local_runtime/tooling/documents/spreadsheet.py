"""Render a .xlsx workbook as one pipe table per sheet, using openpyxl.

The workbook is opened in openpyxl's read-only mode, which streams each sheet's
cells instead of building an object for every cell the file declares. Loading
it whole would let a few bytes of XML expand without bound: a merged range or a
hyperlink is bound to every cell its reference covers, so one ``A1:XFD1048576``
asks for seventeen billion cell objects. Read-only mode leaves out the drawing
parts, so those are read separately with the function openpyxl itself loads
them with. Every sheet is read under both a row/column cap and the shared text
budget.
"""

from __future__ import annotations

from collections.abc import Iterator
from typing import BinaryIO, Final

from openpyxl import load_workbook
from openpyxl.chart._chart import ChartBase
from openpyxl.drawing.image import Image
from openpyxl.drawing.spreadsheet_drawing import SpreadsheetDrawing
from openpyxl.packaging.relationship import get_dependents, get_rels_path
from openpyxl.reader.drawings import find_images
from openpyxl.utils import get_column_letter
from openpyxl.workbook.workbook import Workbook
from openpyxl.worksheet._read_only import ReadOnlyWorksheet

from .document_model import (
    MAX_DOCUMENT_TEXT_CHARS,
    MAX_SHEET_COLUMNS,
    MAX_SHEET_ROWS,
    DocumentChart,
    DocumentImage,
    DocumentTextBudget,
    ExtractedDocument,
    pipe_table,
    resolve_start_unit,
    stored_value_text,
)

_XLSX_NOTES: Final = (
    "Cell formatting, formulas, conditional formatting, data validation, "
    "comments, pivot tables and macros are not included, and chart sheets are "
    "not listed.",
    "Cells carry the values Excel last saved, written as stored rather than as "
    "the sheet displays them. A formula whose result was never saved reads as "
    "an empty cell, and a merged range carries its value in its first cell.",
    f"Each sheet is read up to {MAX_SHEET_ROWS} rows and {MAX_SHEET_COLUMNS} "
    "columns, and trailing empty rows and columns are dropped.",
    "An image is listed by the cell it is anchored at, because the library "
    "does not keep the package part it was loaded from. A chart title taken "
    "from a cell is not resolved.",
)


def extract_xlsx(source: BinaryIO, start_unit: int | None) -> ExtractedDocument:
    # data_only asks for the result Excel stored beside each formula; the
    # formula text itself is not what the model was asked to read.
    workbook = load_workbook(source, read_only=True, data_only=True)
    try:
        return _extract_workbook(workbook, start_unit)
    finally:
        workbook.close()


def _extract_workbook(workbook: Workbook, start_unit: int | None) -> ExtractedDocument:
    sheets = workbook.worksheets
    total_units = len(sheets)
    start = resolve_start_unit(start_unit, total_units=total_units, unit_kind="sheet")
    budget = DocumentTextBudget(
        unit_kind="sheet", total_units=total_units, start_unit=start
    )
    images: list[DocumentImage] = []
    charts: list[DocumentChart] = []
    outline: list[str] = []
    cut_sheets: list[str] = []
    for number, sheet in enumerate(sheets[start - 1 :], start=start):
        name = _sheet_name(sheet)
        outline.append(name)
        sheet_charts, sheet_images = _sheet_drawings(workbook, sheet)
        images.extend(_image_entries(sheet_images, name))
        charts.extend(_chart_entries(sheet_charts, name))
        lines, cut = pipe_table(
            _sheet_rows(sheet),
            max_rows=MAX_SHEET_ROWS,
            max_columns=MAX_SHEET_COLUMNS,
        )
        if cut:
            cut_sheets.append(name)
        # A sheet the budget cannot hold ends this read; total_units still
        # counts every sheet, and next_start_unit says where the rest begin.
        if not budget.add_unit(number, f"## Sheet: {name}", *lines, ""):
            break
    return ExtractedDocument(
        document_format="xlsx",
        unit_kind="sheet",
        total_units=total_units,
        text=budget.finish(),
        content_truncated=budget.stopped or bool(cut_sheets),
        next_start_unit=budget.next_start_unit,
        outline=tuple(outline),
        images=tuple(images),
        charts=tuple(charts),
        notes=(*_XLSX_NOTES, *_cut_note(cut_sheets), *budget.notes()),
    )


def _sheet_name(sheet: ReadOnlyWorksheet) -> str:
    """The sheet's title, marked when Excel does not show the sheet."""

    title = str(sheet.title)
    return title if sheet.sheet_state == "visible" else f"{title} (hidden)"


def _sheet_rows(sheet: ReadOnlyWorksheet) -> Iterator[list[str]]:
    """Cell text row by row, bounded so no sheet is expanded in full.

    The extent a sheet declares is ignored: a writer may omit it or understate
    it, and the stream ends at the sheet's last row anyway. Trailing empty rows
    and columns are dropped. One row and one column past the caps are read so
    that a sheet larger than the caps is reported as cut rather than silently
    narrowed.
    """

    blank_rows = 0
    chars = 0
    for row in sheet.iter_rows(
        max_row=MAX_SHEET_ROWS + 1, max_col=MAX_SHEET_COLUMNS + 1, values_only=True
    ):
        cells = [stored_value_text(value) for value in row]
        while cells and not cells[-1]:
            cells.pop()
        if not cells:
            blank_rows += 1
            continue
        for _ in range(blank_rows):
            yield []
        blank_rows = 0
        yield cells
        # Rendering only adds to these characters, so stopping here guarantees
        # the text budget rejects a line and reports the document as cut.
        chars += sum(len(cell) for cell in cells)
        if chars >= MAX_DOCUMENT_TEXT_CHARS:
            return


def _sheet_drawings(
    workbook: Workbook, sheet: ReadOnlyWorksheet
) -> tuple[list[ChartBase], list[Image]]:
    """The charts and pictures on one sheet, as openpyxl loads a whole workbook.

    This is the drawing step of openpyxl's own worksheet reader, which read-only
    mode skips. openpyxl keeps no public accessor for the package or for the
    part a read-only sheet streams from.
    """

    archive = workbook._archive
    relationships_path = get_rels_path(sheet._worksheet_path)
    if relationships_path not in archive.namelist():
        return [], []
    charts: list[ChartBase] = []
    images: list[Image] = []
    relationships = get_dependents(archive, relationships_path)
    for drawing in relationships.find(SpreadsheetDrawing._rel_type):
        drawing_charts, drawing_images = find_images(archive, drawing.target)
        charts.extend(drawing_charts)
        images.extend(drawing_images)
    return charts, images


def _image_entries(images: list[Image], name: str) -> Iterator[DocumentImage]:
    """Pictures on this sheet, counted per anchor cell."""

    counts: dict[str, int] = {}
    for image in images:
        cell = _anchor_cell(image.anchor)
        counts[cell] = counts.get(cell, 0) + 1
    for cell, count in counts.items():
        yield DocumentImage(location=name, ref=cell, count=count)


def _chart_entries(charts: list[ChartBase], name: str) -> Iterator[DocumentChart]:
    for chart in charts:
        cell = _anchor_cell(chart.anchor)
        yield DocumentChart(
            location=f"{name}!{cell}" if cell else name,
            title=_chart_title(chart.title),
        )


def _anchor_cell(anchor: object) -> str:
    """The cell a drawing hangs from, empty for one pinned to the page itself."""

    marker = getattr(anchor, "_from", None)
    if marker is None:
        return ""
    return f"{get_column_letter(int(marker.col) + 1)}{int(marker.row) + 1}"


def _chart_title(title: object) -> str:
    """The title text a chart carries, empty when it points at a cell instead."""

    rich = getattr(getattr(title, "tx", None), "rich", None)
    if rich is None:
        return ""
    return "".join(
        str(run.t)
        for paragraph in rich.p or ()
        for run in paragraph.r or ()
        if run.t is not None
    )


def _cut_note(cut_sheets: list[str]) -> tuple[str, ...]:
    if not cut_sheets:
        return ()
    return (
        f"These sheets are larger than {MAX_SHEET_ROWS} rows by "
        f"{MAX_SHEET_COLUMNS} columns and were cut: {', '.join(cut_sheets)}.",
    )


__all__ = ["extract_xlsx"]
