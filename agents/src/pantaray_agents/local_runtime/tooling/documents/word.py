"""Render a .docx body as headings, lists and tables, using python-docx."""

from __future__ import annotations

import re
from typing import BinaryIO, Final

from docx import Document
from docx.document import Document as DocxDocument
from docx.oxml.ns import qn
from docx.oxml.xmlchemy import BaseOxmlElement
from docx.table import Table
from docx.text.paragraph import Paragraph

from .document_model import (
    DocumentImage,
    DocumentTextBudget,
    ExtractedDocument,
    pipe_table,
    resolve_start_unit,
)

_HEADING_STYLE: Final = re.compile(r"heading\s*([1-9])", re.IGNORECASE)
_LIST_STYLE: Final = re.compile(r"list", re.IGNORECASE)
# Each note stands alone and stays short, because the read tool's summary caps
# how long a single note may be and silently shortens one that runs past it.
_DOCX_NOTES: Final = (
    "Formatting, comments, tracked changes, headers, footers, footnotes and "
    "endnotes are not included. Charts and other drawings that are not "
    "embedded pictures are not listed.",
    "List numbering is not resolved, so a numbered item carries the same "
    "marker as a bullet.",
    "A page count is unavailable without rendering the document.",
)


def extract_docx(source: BinaryIO, start_unit: int | None) -> ExtractedDocument:
    document = Document(source)
    total_units = len(document.paragraphs)
    start = resolve_start_unit(
        start_unit, total_units=total_units, unit_kind="paragraph"
    )
    budget = DocumentTextBudget(
        unit_kind="paragraph", total_units=total_units, start_unit=start
    )
    images: list[DocumentImage] = []
    outline: list[str] = []
    paragraphs = 0
    tables = 0
    table_cut = False
    for block in document.element.body.iterchildren():
        is_paragraph = block.tag == qn("w:p")
        if is_paragraph:
            paragraphs += 1
        elif block.tag == qn("w:tbl"):
            tables += 1
        else:
            continue
        # A table sits between paragraphs, so it is read as part of the
        # paragraph it follows: a read resuming at a later paragraph skips it
        # because the read that reached that paragraph already carried it.
        number = max(paragraphs, 1)
        if number < start:
            continue
        if is_paragraph:
            location = f"paragraph {paragraphs}"
            line = _paragraph_line(Paragraph(block, document))
            if line.startswith("#"):
                outline.append(line)
            lines = [line]
        else:
            location = f"table {tables}"
            lines, cut = pipe_table(
                [cell.text for cell in row.cells] for row in Table(block, document).rows
            )
            table_cut = table_cut or cut
        images.extend(_block_images(block, document, location))
        if not budget.add_unit(number, *lines):
            break
    return ExtractedDocument(
        document_format="docx",
        unit_kind="paragraph",
        total_units=total_units,
        text=budget.finish(),
        content_truncated=budget.stopped or table_cut,
        next_start_unit=budget.next_start_unit,
        outline=tuple(outline),
        images=tuple(images),
        notes=(*_DOCX_NOTES, *budget.notes()),
    )


def _paragraph_line(paragraph: Paragraph) -> str:
    text = paragraph.text
    if not text.strip():
        return ""
    style = (paragraph.style.name or "") if paragraph.style is not None else ""
    heading = _HEADING_STYLE.fullmatch(style)
    if heading is not None:
        return f"{'#' * int(heading.group(1))} {text}"
    if style.casefold() == "title":
        return f"# {text}"
    if _LIST_STYLE.search(style):
        return f"- {text}"
    return text


def _block_images(
    block: BaseOxmlElement, document: DocxDocument, location: str
) -> list[DocumentImage]:
    """Images this block displays, named by the package part they live in."""

    relationships = document.part.rels
    counts: dict[str, int] = {}
    for blip in block.findall(f".//{qn('a:blip')}"):
        reference = blip.get(qn("r:embed"))
        if reference is None or reference not in relationships:
            continue
        relationship = relationships[reference]
        if relationship.is_external:
            continue
        part_name = str(relationship.target_part.partname).lstrip("/")
        counts[part_name] = counts.get(part_name, 0) + 1
    return [
        DocumentImage(location=location, ref=part_name, count=count)
        for part_name, count in counts.items()
    ]


__all__ = ["extract_docx"]
